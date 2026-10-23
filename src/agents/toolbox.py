"""Per-role tools for the three FP&A agents, and the checks that keep each agent inside its contract.

Each role gets the scripts listed under its `allowed_scripts` in `config/workspace_governance.json`,
a scoped `read_file`, and one tool to hand off its result. Three checks run here, outside the model:

- Read scope: `read_file` and `read_workbook` open only the context files and the role's read/write scopes.
- Write scope: every script call hashes `data/` and `output/` before and after, and any changed file
  outside the role's `write_scopes` is recorded as a violation.
- Citations: a hand-off is refused if it states a figure that never appeared in the role's tool results.

This module does not import the Agent SDK, so the tests run it with no API key.
"""

from __future__ import annotations

from dataclasses import dataclass
import fnmatch
import hashlib
import json
from pathlib import Path
import re
from typing import Any, Callable, Dict, List, Optional, Set, Tuple

import openpyxl
from openpyxl.utils.cell import range_boundaries

from src.build_rough_budget import run_stage_3
from src.build_trends import run_stage_2
from src.cut_review_rules import run_stage_4
from src.domain import AUDIT_REPORT_PATH, CONFIG_PATH, ROOT_DIR, WORKBOOK_PATH
from src.prepare_data import run_stage_1_on_existing_raw
from src.verify_governance import run_governance_audit


ROLES: Tuple[str, ...] = ("variance_analyst", "budget_architect", "governance_auditor")
AUDITOR = "governance_auditor"
ROLE_TITLES: Dict[str, str] = {
    "variance_analyst": "Variance Analyst",
    "budget_architect": "Budget Architect",
    "governance_auditor": "Governance Auditor",
}
CONTEXT_FILES: Tuple[str, ...] = (
    "workspace/GLOBAL_CONTEXT.md",
    "workspace/WORKSPACE_CONTEXT.md",
    "config/workspace_governance.json",
)
GUARDED_DIRS: Tuple[str, ...] = ("data", "output")
MAX_READ_CHARS = 20_000
MAX_WORKBOOK_CELLS = 400
FIGURE_PATTERN = re.compile(r"\d[\d,]*\.\d+")


def _run_verify_governance() -> Dict[str, Any]:
    """Run Stage 5 and return the report whether it passes or fails."""
    try:
        return {"status": "PASSED", "report": run_governance_audit().to_dict()}
    except RuntimeError as exc:
        if not str(exc).startswith("Governance verification failed"):
            raise
        return {"status": "FAILED", "report": json.loads(AUDIT_REPORT_PATH.read_text(encoding="utf-8"))}


# Script path in the policy file -> (tool name, description, entry point).
SCRIPT_TOOLS: Dict[str, Tuple[str, str, Callable[[], Any]]] = {
    "src/prepare_data.py": (
        "run_prepare_data",
        "Stage 1. Validate the ledger extract in data/raw/ and split one-off spikes from recurring actuals (ADR-0001).",
        run_stage_1_on_existing_raw,
    ),
    "src/build_trends.py": (
        "run_build_trends",
        "Stage 2. Compare TTM with P12M on recurring amounts and write ttm_trend_metrics.csv.",
        run_stage_2,
    ),
    "src/build_rough_budget.py": (
        "run_build_rough_budget",
        "Stage 3. Build output/FY2027_Rolling_Budget_Model.xlsx from the trend metrics, with live roll-up formulas.",
        run_stage_3,
    ),
    "src/cut_review_rules.py": (
        "run_cut_review_rules",
        "Stage 4. Apply ADR-0003 cuts, write the controller locks, re-run with +1.50 points and write the lock proof.",
        run_stage_4,
    ),
    "src/verify_governance.py": (
        "run_verify_governance",
        "Stage 5. Read the workbook back and check ADR-0001..0004. Returns status PASSED or FAILED and the report.",
        _run_verify_governance,
    ),
}


@dataclass(frozen=True)
class RoleScope:
    """One role's access, read from `config/workspace_governance.json`."""

    role: str
    contract_path: str
    allowed_scripts: Tuple[str, ...]
    read_scopes: Tuple[str, ...]
    write_scopes: Tuple[str, ...]

    def can_read(self, rel_path: str) -> bool:
        patterns = CONTEXT_FILES + (self.contract_path,) + self.read_scopes + self.write_scopes
        return any(fnmatch.fnmatchcase(rel_path, p) for p in patterns)

    def can_write(self, rel_path: str) -> bool:
        return any(fnmatch.fnmatchcase(rel_path, p) for p in self.write_scopes)


def load_role_scopes(config_path: Path = CONFIG_PATH) -> Dict[str, RoleScope]:
    """Load the three role scopes from the policy file."""
    config = json.loads(config_path.read_text(encoding="utf-8"))
    tier_3 = config["context_hierarchy"]["tier_3_agents"]
    return {
        role: RoleScope(
            role=role,
            contract_path=tier_3[role]["contract_path"],
            allowed_scripts=tuple(tier_3[role]["allowed_scripts"]),
            read_scopes=tuple(tier_3[role]["read_scopes"]),
            write_scopes=tuple(tier_3[role]["write_scopes"]),
        )
        for role in ROLES
    }


def build_system_prompt(role: str) -> str:
    """Return the runtime preamble followed by the three context tiers, in load order."""
    scope = load_role_scopes()[role]
    if role == AUDITOR:
        finish = (
            "Finish by calling submit_verdict once. Approve only if run_verify_governance returned PASSED; the tool "
            "refuses an approval otherwise. If anything failed, reject and send the work back to the agent that "
            "executes the failing rule, as named on the ADR's 'Owner Agent' line."
        )
    else:
        finish = (
            "Finish by calling submit_handoff once. Use status \"blocked\" if any completion criterion in your "
            "contract failed, and say which one."
        )
    preamble = "\n".join(
        [
            f"You are the {ROLE_TITLES[role]} agent in the FY2027 rolling-budget workspace, running under the "
            "Claude Agent SDK. The three files below are your instructions, loaded in order: global rules, this "
            "budget cycle, then your role contract.",
            "",
            "How this run works:",
            "- Your tools are the only way to act. Where your contract says to execute `python3 src/<script>.py`, "
            "call the matching run_ tool. You cannot run shell commands or edit files.",
            "- read_file opens only paths in your access scope. If a read is refused, do not try another route.",
            "- Do no arithmetic. Copy every figure from a tool result. The hand-off tool refuses figures that do "
            "not appear in your tool results or your task.",
            "- Check each completion criterion in your contract against the tool output.",
            f"- {finish}",
            "- Keep the hand-off short: what you ran, the figures the next agent needs, and anything that failed.",
        ]
    )
    tiers = [
        ("Tier 1", "workspace/GLOBAL_CONTEXT.md"),
        ("Tier 2", "workspace/WORKSPACE_CONTEXT.md"),
        ("Tier 3", scope.contract_path),
    ]
    sections = [preamble]
    for label, rel_path in tiers:
        text = (ROOT_DIR / rel_path).read_text(encoding="utf-8")
        sections.append(f"<context tier=\"{label}\" path=\"{rel_path}\">\n{text.strip()}\n</context>")
    return "\n\n".join(sections)


def figures_in(text: str) -> Set[float]:
    """Return the decimal figures in `text` as unsigned floats ("12,877.22" and "-12877.22" both give 12877.22)."""
    return {round(abs(float(m.replace(",", ""))), 6) for m in FIGURE_PATTERN.findall(text)}


def snapshot_outputs() -> Dict[str, str]:
    """Hash every file under the guarded directories."""
    hashes: Dict[str, str] = {}
    for dir_name in GUARDED_DIRS:
        base = ROOT_DIR / dir_name
        if not base.exists():
            continue
        for path in sorted(base.rglob("*")):
            if path.is_file():
                hashes[path.relative_to(ROOT_DIR).as_posix()] = hashlib.sha256(path.read_bytes()).hexdigest()
    return hashes


def changed_paths(before: Dict[str, str], after: Dict[str, str]) -> List[str]:
    """Return files created, changed or deleted between two snapshots."""
    return sorted(p for p in before.keys() | after.keys() if before.get(p) != after.get(p))


@dataclass(frozen=True)
class ToolSpec:
    """A tool as the model sees it, independent of the SDK."""

    name: str
    description: str
    input_schema: Dict[str, Any]
    handler: Callable[[Dict[str, Any]], str]


class RoleToolbox:
    """The tools for one agent run, plus what that run handed off."""

    def __init__(
        self,
        role: str,
        task: str,
        record: Callable[[Dict[str, Any]], None],
        script_tools: Optional[Dict[str, Tuple[str, str, Callable[[], Any]]]] = None,
    ) -> None:
        self.role = role
        self.scope = load_role_scopes()[role]
        self.record = record
        self.script_tools = script_tools if script_tools is not None else SCRIPT_TOOLS
        self.sources: List[str] = [task]
        self.write_violations: List[str] = []
        self.last_audit_status: Optional[str] = None
        self.result: Optional[Dict[str, Any]] = None
        self.specs: Dict[str, ToolSpec] = {spec.name: spec for spec in self._build_specs()}

    def call(self, name: str, args: Dict[str, Any]) -> str:
        """Run one tool, record the call and its output, and return the output."""
        spec = self.specs.get(name)
        output = spec.handler(args) if spec else f"ERROR: {self.role} has no tool named {name}."
        self.record({"type": "tool_call", "tool": name, "input": args, "output": output})
        return output

    def _build_specs(self) -> List[ToolSpec]:
        specs: List[ToolSpec] = []
        for script in self.scope.allowed_scripts:
            tool_name, description, entry_point = self.script_tools[script]
            specs.append(
                ToolSpec(
                    name=tool_name,
                    description=description,
                    input_schema={"type": "object", "properties": {}},
                    handler=lambda _args, s=script, fn=entry_point: self._run_script(s, fn),
                )
            )
        specs.append(
            ToolSpec(
                name="read_file",
                description="Read a text file (Markdown, CSV, JSON) inside your access scope. Paths are relative to the repo root.",
                input_schema={"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]},
                handler=self._read_file,
            )
        )
        if self.scope.can_read(WORKBOOK_PATH.relative_to(ROOT_DIR).as_posix()):
            specs.append(
                ToolSpec(
                    name="read_workbook",
                    description=(
                        "Read cells from output/FY2027_Rolling_Budget_Model.xlsx. Formulas come back as written "
                        f"(=SUM(...)). Sheets: FY2027_Budget_Model, Executive_Summary, ADR_Audit_Trail. "
                        f"At most {MAX_WORKBOOK_CELLS} cells per call."
                    ),
                    input_schema={
                        "type": "object",
                        "properties": {"sheet": {"type": "string"}, "cell_range": {"type": "string"}},
                        "required": ["sheet", "cell_range"],
                    },
                    handler=self._read_workbook,
                )
            )
        if self.role == AUDITOR:
            specs.append(
                ToolSpec(
                    name="submit_verdict",
                    description="Approve the run, or reject it and name the agent that must redo its stages.",
                    input_schema={
                        "type": "object",
                        "properties": {
                            "verdict": {"type": "string", "enum": ["approve", "reject"]},
                            "send_back_to": {"type": "string", "enum": ["none", "variance_analyst", "budget_architect"]},
                            "reasons": {"type": "string"},
                        },
                        "required": ["verdict", "send_back_to", "reasons"],
                    },
                    handler=self._submit_verdict,
                )
            )
        else:
            specs.append(
                ToolSpec(
                    name="submit_handoff",
                    description="Hand your result to the next agent. Call once, at the end.",
                    input_schema={
                        "type": "object",
                        "properties": {
                            "status": {"type": "string", "enum": ["complete", "blocked"]},
                            "summary": {"type": "string"},
                        },
                        "required": ["status", "summary"],
                    },
                    handler=self._submit_handoff,
                )
            )
        return specs

    def _run_script(self, script: str, entry_point: Callable[[], Any]) -> str:
        before = snapshot_outputs()
        try:
            output: Any = entry_point()
            status = "ok"
        except Exception as exc:  # the agent reports the failure; the coordinator decides what happens next
            output = f"{type(exc).__name__}: {exc}"
            status = "error"
        written = changed_paths(before, snapshot_outputs())
        violations = [p for p in written if not self.scope.can_write(p)]
        self.write_violations.extend(violations)
        if isinstance(output, dict) and "status" in output and "report" in output:
            self.last_audit_status = output["status"]
        text = json.dumps(
            {
                "script": script,
                "status": status,
                "output": output,
                "files_written": written,
                "write_scope_violations": violations,
            },
            indent=2,
        )
        self.sources.append(text)
        return text

    def _resolve(self, path: str) -> Optional[str]:
        """Return `path` relative to the repo root, or None if it points outside it."""
        root = ROOT_DIR.resolve()
        resolved = (root / path).resolve()
        if not resolved.is_relative_to(root):
            return None
        return resolved.relative_to(root).as_posix()

    def _read_file(self, args: Dict[str, Any]) -> str:
        rel_path = self._resolve(str(args.get("path", "")))
        if rel_path is None or not self.scope.can_read(rel_path):
            return f"REFUSED: {args.get('path')} is outside the {self.role} read scope."
        if rel_path.endswith(".xlsx"):
            return "REFUSED: use read_workbook for the workbook."
        path = ROOT_DIR / rel_path
        if not path.is_file():
            return f"ERROR: {rel_path} does not exist."
        text = path.read_text(encoding="utf-8")
        if len(text) > MAX_READ_CHARS:
            text = text[:MAX_READ_CHARS] + f"\n[truncated: showing the first {MAX_READ_CHARS} of {len(text)} characters]"
        self.sources.append(text)
        return text

    def _read_workbook(self, args: Dict[str, Any]) -> str:
        path = WORKBOOK_PATH
        if not path.is_file():
            return "ERROR: the workbook does not exist yet."
        try:
            min_col, min_row, max_col, max_row = range_boundaries(str(args.get("cell_range", "")))
        except ValueError as exc:
            return f"ERROR: {exc}"
        if (max_col - min_col + 1) * (max_row - min_row + 1) > MAX_WORKBOOK_CELLS:
            return f"ERROR: range is larger than {MAX_WORKBOOK_CELLS} cells."
        wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
        try:
            sheet = str(args.get("sheet", ""))
            if sheet not in wb.sheetnames:
                return f"ERROR: no sheet {sheet}. Sheets: {', '.join(wb.sheetnames)}"
            lines = [
                " | ".join(f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None)
                for row in wb[sheet].iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
            ]
        finally:
            wb.close()
        text = "\n".join(line for line in lines if line) or "(empty range)"
        self.sources.append(text)
        return text

    def _citation_refusal(self, text: str) -> Optional[str]:
        """Return a refusal if `text` states a figure that is not in the task or any tool result."""
        cited: Set[float] = set()
        for source in self.sources:
            cited |= figures_in(source)
        uncited = sorted(figures_in(text) - cited)
        if not uncited:
            return None
        return (
            f"REFUSED: these figures are not in any tool result you received: {uncited}. "
            "Copy figures exactly from tool output, or leave them out."
        )

    def _submit_handoff(self, args: Dict[str, Any]) -> str:
        if self.result is not None:
            return "REFUSED: you already handed off."
        summary = str(args.get("summary", ""))
        refusal = self._citation_refusal(summary)
        if refusal:
            return refusal
        self.result = {"status": args.get("status"), "summary": summary}
        return "Hand-off accepted. Stop here."

    def _submit_verdict(self, args: Dict[str, Any]) -> str:
        if self.result is not None:
            return "REFUSED: you already submitted a verdict."
        verdict = args.get("verdict")
        send_back_to = args.get("send_back_to")
        reasons = str(args.get("reasons", ""))
        if verdict == "approve" and self.last_audit_status != "PASSED":
            return "REFUSED: you can approve only after run_verify_governance returns PASSED in this session."
        if verdict == "reject" and send_back_to not in ("variance_analyst", "budget_architect"):
            return "REFUSED: a rejection must name variance_analyst or budget_architect in send_back_to."
        refusal = self._citation_refusal(reasons)
        if refusal:
            return refusal
        self.result = {"verdict": verdict, "send_back_to": send_back_to, "reasons": reasons}
        return "Verdict recorded. Stop here."
