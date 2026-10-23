"""Coordinator: passes work from the variance analyst to the budget architect to the governance auditor.

The auditor either approves, or rejects and names the agent that must redo its stages. Work restarts at
that agent and runs forward again. The coordinator allows `max_rejections` send-backs and stops at the next rejection.

The coordinator takes a runner, `runner(role, task, toolbox) -> stats`. `sdk_runner.make_sdk_runner`
runs a real agent; `ReplayRunner` re-issues the tool calls from a recording, which is what the tests use.
"""

from __future__ import annotations

from dataclasses import dataclass, field
import json
from pathlib import Path
import shutil
from typing import Any, Callable, Dict, List, Optional

import openpyxl

from src.agents.toolbox import AUDITOR, ROLE_TITLES, ROLES, RoleToolbox
from src.domain import BUDGET_MODEL_HEADERS, OUTPUT_DIR, PROCESSED_DIR, ROOT_DIR, WORKBOOK_PATH, BudgetLineItem
from src.prepare_data import generate_raw_datasets


AgentRunner = Callable[[str, str, RoleToolbox], Dict[str, Any]]

TASKS: Dict[str, str] = {
    "variance_analyst": (
        "Run Stage 1 and Stage 2 for the FY2027 cycle. The ledger extract is already in data/raw/. "
        "Check both completion criteria in your contract against the tool output, then hand off."
    ),
    "budget_architect": (
        "Run Stage 3 and Stage 4 for the FY2027 cycle, starting from the variance analyst's trend metrics. "
        "Check both completion criteria in your contract against the tool output, then hand off."
    ),
    "governance_auditor": (
        "Run Stage 5 on the finished workbook and decide whether this budget run is approved."
    ),
}

RECORDED_OUTPUTS = (
    "data/processed/excluded_one_off_items.csv",
    "data/processed/ttm_trend_metrics.csv",
    "data/processed/rerun_preservation_proof.json",
    "data/processed/governance_audit_report.json",
    "output/FY2027_Rolling_Budget_Model.xlsx",
)

TAMPER_COST_CENTER = "CC-208"
TAMPER_DELTA_NIU = 5.00


@dataclass
class Recorder:
    """Collects every event of a run in order."""

    events: List[Dict[str, Any]] = field(default_factory=list)
    run: int = 0
    role: str = "coordinator"

    def log(self, event: Dict[str, Any]) -> None:
        self.events.append({"seq": len(self.events) + 1, "run": self.run, "role": self.role, **event})


def build_task(role: str, handoffs: Dict[str, Dict[str, Any]], rejection: Optional[Dict[str, Any]]) -> str:
    """Write the prompt for one agent run: its task, the upstream hand-offs, and any rejection to fix."""
    parts = [TASKS[role]]
    for upstream in ROLES[: ROLES.index(role)]:
        if upstream in handoffs:
            parts.append(f"Hand-off from the {ROLE_TITLES[upstream]} ({handoffs[upstream]['status']}):\n{handoffs[upstream]['summary']}")
    if rejection is not None:
        parts.append(
            "The Governance Auditor rejected the previous result and sent it back to you. Its reasons:\n"
            f"{rejection['reasons']}\nFix the cause by re-running your stages, then hand off again."
        )
    return "\n\n".join(parts)


def tamper_locked_row(cost_center: str = TAMPER_COST_CENTER, delta_niu: float = TAMPER_DELTA_NIU) -> Dict[str, Any]:
    """Change one locked row's final budget in the saved workbook, as an unsigned hand edit would."""
    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    ws = wb["FY2027_Budget_Model"]
    final_col = BUDGET_MODEL_HEADERS.index("Final_Budget_NIU") + 1
    for row_idx in range(2, ws.max_row + 1):
        item = BudgetLineItem.from_worksheet_row(ws, row_idx)
        if item.cost_center == cost_center:
            if not item.is_locked:
                raise ValueError(f"{cost_center} is not a locked row")
            after = round(item.final_budget_niu + delta_niu, 2)
            cell = ws.cell(row=row_idx, column=final_col, value=after)
            wb.save(WORKBOOK_PATH)
            return {"cost_center": cost_center, "cell": cell.coordinate, "before": item.final_budget_niu, "after": after}
    raise ValueError(f"{cost_center} not found in the workbook")


def run_pipeline(
    runner: AgentRunner,
    recorder: Recorder,
    tamper: bool = False,
    max_rejections: int = 2,
) -> Dict[str, Any]:
    """Run the three agents in order until the auditor approves, an agent blocks, or rejections run out."""
    for stale_dir in (PROCESSED_DIR, OUTPUT_DIR):
        shutil.rmtree(stale_dir, ignore_errors=True)
    generate_raw_datasets()
    recorder.log({"type": "coordinator", "text": "Wrote the synthetic ledger extract to data/raw/. No agent can write there."})

    handoffs: Dict[str, Dict[str, Any]] = {}
    rejection: Optional[Dict[str, Any]] = None
    rejections = 0
    tampered = False
    runs: List[Dict[str, Any]] = []
    idx = 0

    while True:
        role = ROLES[idx]
        task = build_task(role, handoffs, rejection)
        recorder.run += 1
        recorder.role = role
        recorder.log({"type": "task", "text": task})
        toolbox = RoleToolbox(role, task, recorder.log)
        stats = runner(role, task, toolbox)
        runs.append({"run": recorder.run, "role": role, "result": toolbox.result, **stats})
        recorder.log({"type": "result", "result": toolbox.result})
        recorder.role = "coordinator"

        if toolbox.write_violations:
            outcome = f"halted: {role} wrote outside its scope: {toolbox.write_violations}"
            break
        if toolbox.result is None:
            outcome = f"halted: {role} finished without handing off"
            break

        if role != AUDITOR:
            if toolbox.result["status"] != "complete":
                outcome = f"halted: {role} reported blocked"
                break
            handoffs[role] = toolbox.result
            rejection = None
            idx += 1
            if tamper and role == "budget_architect" and not tampered:
                change = tamper_locked_row()
                tampered = True
                recorder.log(
                    {
                        "type": "coordinator",
                        "text": (
                            f"Fault injected for the demo: changed {change['cost_center']} Final_Budget_NIU "
                            f"(cell {change['cell']}) from {change['before']:.2f} to {change['after']:.2f}. "
                            "The row is LOCKED, and this edit is not in the signed overrides file."
                        ),
                    }
                )
            continue

        if toolbox.result["verdict"] == "approve":
            outcome = "approved"
            break
        rejections += 1
        if rejections > max_rejections:
            outcome = f"halted: {rejections} rejections"
            break
        rejection = toolbox.result
        idx = ROLES.index(toolbox.result["send_back_to"])
        for later in ROLES[idx:]:
            handoffs.pop(later, None)
        recorder.log({"type": "coordinator", "text": f"Rejected. Sending the work back to the {ROLE_TITLES[ROLES[idx]]}."})

    recorder.log({"type": "coordinator", "text": f"Outcome: {outcome}."})
    return {"outcome": outcome, "tamper": tamper, "rejections": rejections, "runs": runs}


def save_recording(out_dir: Path, recorder: Recorder, summary: Dict[str, Any]) -> None:
    """Write transcript.jsonl, transcript.md, summary.json and the run's outputs to `out_dir`."""
    out_dir.mkdir(parents=True, exist_ok=True)
    with (out_dir / "transcript.jsonl").open("w", encoding="utf-8") as fh:
        for event in recorder.events:
            fh.write(json.dumps(event) + "\n")
    (out_dir / "summary.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    (out_dir / "transcript.md").write_text(render_transcript(recorder.events, summary), encoding="utf-8")
    outputs_dir = out_dir / "outputs"
    shutil.rmtree(outputs_dir, ignore_errors=True)
    outputs_dir.mkdir()
    for rel_path in RECORDED_OUTPUTS:
        source = ROOT_DIR / rel_path
        if source.exists():
            shutil.copy2(source, outputs_dir / source.name)


def load_recording(out_dir: Path) -> List[Dict[str, Any]]:
    lines = (out_dir / "transcript.jsonl").read_text(encoding="utf-8").splitlines()
    return [json.loads(line) for line in lines if line.strip()]


def render_transcript(events: List[Dict[str, Any]], summary: Dict[str, Any]) -> str:
    """Render the events as Markdown, with long tool output folded."""
    lines = [
        f"# Recorded run: {summary['outcome']}",
        "",
        f"Model: `{summary.get('model', 'unknown')}`. Agent runs: {len(summary['runs'])}. "
        f"Rejections: {summary['rejections']}. Fault injected: {'yes' if summary['tamper'] else 'no'}.",
        "",
        "Generated by `src/run_agents.py`. The full event log is `transcript.jsonl`.",
    ]
    for event in events:
        kind = event["type"]
        if kind == "coordinator":
            lines += ["", f"> **Coordinator:** {event['text']}"]
        elif kind == "task":
            lines += ["", f"## Run {event['run']}: {ROLE_TITLES[event['role']]}", "", "**Task**", "", _quote(event["text"])]
        elif kind == "assistant_text":
            lines += ["", event["text"]]
        elif kind == "tool_call":
            args = json.dumps(event["input"]) if event["input"] else ""
            lines += [
                "",
                f"<details><summary><code>{event['tool']}({args})</code></summary>",
                "",
                "```",
                event["output"].strip(),
                "```",
                "",
                "</details>",
            ]
        elif kind == "result" and event["result"] is not None:
            lines += ["", "**Result**", "", "```json", json.dumps(event["result"], indent=2), "```"]
    return "\n".join(lines) + "\n"


def _quote(text: str) -> str:
    return "\n".join(f"> {line}" if line else ">" for line in text.splitlines())


class ReplayRunner:
    """Re-issues a recording's tool calls against the real tools, and notes any output that differs.

    The model's text is not replayed. What is checked is that each tool, given the recorded input,
    still returns the recorded output, and that the hand-offs drive the coordinator the same way.
    """

    def __init__(self, events: List[Dict[str, Any]]) -> None:
        self.calls_by_run: Dict[int, List[Dict[str, Any]]] = {}
        self.roles_by_run: Dict[int, str] = {}
        for event in events:
            if event["type"] == "tool_call":
                self.calls_by_run.setdefault(event["run"], []).append(event)
            if event["type"] == "task":
                self.roles_by_run[event["run"]] = event["role"]
        self.next_run = 1
        self.mismatches: List[str] = []

    def __call__(self, role: str, task: str, toolbox: RoleToolbox) -> Dict[str, Any]:
        run = self.next_run
        self.next_run += 1
        if self.roles_by_run.get(run) != role:
            self.mismatches.append(f"run {run}: recording has {self.roles_by_run.get(run)}, coordinator asked for {role}")
            return {}
        for call in self.calls_by_run.get(run, []):
            output = toolbox.call(call["tool"], call["input"])
            if output != call["output"]:
                self.mismatches.append(f"run {run} {role}.{call['tool']}: output differs from the recording")
        return {}
